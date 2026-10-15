# -*- coding: utf-8 -*-
"""Genera la carta de solicitud a Osinergmin y la guía para la reunión.
    python generar_carta_y_guia_osinergmin.py
Los campos entre corchetes [ ] se completan a mano antes de enviar."""
import os
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
FUENTE = "Arial"


def base(margen=2.5):
    d = Document()
    for s in d.sections:
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin = s.right_margin = Cm(margen)
        s.top_margin = s.bottom_margin = Cm(2.2)
    st = d.styles["Normal"]
    st.font.name = FUENTE
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15
    return d


def p(d, texto="", negrita=False, alinear=None, tam=None, antes=None, despues=None, sangria=None, cursiva=False):
    par = d.add_paragraph()
    for i, trozo in enumerate(texto.split("**")):
        if trozo:
            r = par.add_run(trozo)
            r.bold = negrita or (i % 2 == 1)
            r.italic = cursiva
            if tam:
                r.font.size = Pt(tam)
    if alinear is not None:
        par.alignment = alinear
    if antes is not None:
        par.paragraph_format.space_before = Pt(antes)
    if despues is not None:
        par.paragraph_format.space_after = Pt(despues)
    if sangria is not None:
        par.paragraph_format.left_indent = Cm(sangria)
    return par


def lista(d, items, numerada=False, sangria=0.6):
    for k, t in enumerate(items, 1):
        par = p(d, ("%d. " % k if numerada else "– ") + t, alinear=WD_ALIGN_PARAGRAPH.JUSTIFY, despues=3)
        par.paragraph_format.left_indent = Cm(sangria + 0.5)
        par.paragraph_format.first_line_indent = Cm(-0.5)


def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def tabla(d, filas, anchos_cm, encabezado="1F2937"):
    t = d.add_table(rows=len(filas), cols=len(anchos_cm))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, fila in enumerate(filas):
        for j, v in enumerate(fila):
            c = t.cell(i, j)
            c.width = Cm(anchos_cm[j])
            c.text = ""
            par = c.paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            for k, trozo in enumerate(str(v).split("**")):
                if trozo:
                    r = par.add_run(trozo)
                    r.font.size = Pt(9.5)
                    r.bold = (i == 0) or (k % 2 == 1)
                    if i == 0:
                        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if i == 0:
                sombrear(c, encabezado)
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def titulo(d, t, nivel=1):
    par = p(d, t, negrita=True, tam=13 if nivel == 1 else 11.5, antes=10 if nivel == 1 else 6, despues=4)
    for r in par.runs:
        r.font.color.rgb = RGBColor(0x0F, 0x5E, 0x5C)
    par.paragraph_format.keep_with_next = True
    return par


J = WD_ALIGN_PARAGRAPH.JUSTIFY

# ================================================================== CARTA
c = base()
p(c, "SUMILLA: Solicitamos conformidad para el uso académico de información del SCOP, designación de "
     "especialistas validadores y reunión de presentación de resultados.", negrita=True, alinear=WD_ALIGN_PARAGRAPH.RIGHT,
  tam=10, despues=14).paragraph_format.left_indent = Cm(7)
p(c, "Lima, [día] de octubre de 2026", alinear=WD_ALIGN_PARAGRAPH.RIGHT, despues=12)
p(c, "Señor(a):", despues=0)
p(c, "[Nombres y apellidos]", negrita=True, despues=0)
p(c, "Jefe(a) de la División de Supervisión Regional", despues=0)
p(c, "Gerencia de Supervisión de Energía", despues=0)
p(c, "Organismo Supervisor de la Inversión en Energía y Minería – Osinergmin", despues=0)
p(c, "Presente.-", despues=10)
par = p(c, "Referencia:\tExp. SIGED 202600212189 – Informe GSE/DSR-1875-2026, del 23 de septiembre de 2026", despues=12)
par.paragraph_format.left_indent = Cm(2.6); par.paragraph_format.first_line_indent = Cm(-2.6)
p(c, "De nuestra consideración:", despues=8)

p(c, "Es grato dirigirnos a usted para saludarle cordialmente. Somos David Edison Guerra Ordoñez, "
     "Richard [nombre como figura en el DNI] Yanqui Inocente y Kevin Javier Tello Rivas, estudiantes de la Escuela Profesional "
     "de Ingeniería de Sistemas de la Universidad Peruana Unión, identificados con DNI N.° [ ], [ ] y [ ], "
     "respectivamente, y desarrollamos la tesis titulada «Modelo de clasificación supervisada para la priorización de "
     "establecimientos de venta de combustibles según su probabilidad de variación de precios», bajo la asesoría de "
     "[grado, nombres y apellidos del asesor].", alinear=J)
p(c, "Mediante el documento de la referencia, su representada atendió nuestra solicitud de acceso a la información pública "
     "e indicó que los precios de los combustibles se encuentran publicados en el portal institucional, conforme al "
     "artículo 17 del Procedimiento de entrega de información sobre Precios del Mercado Interno de Combustibles Derivados "
     "de Hidrocarburos (PRICE), aprobado por Resolución de Consejo Directivo N.° 256-2021-OS/CD. Con esa información, "
     "correspondiente al Registro de Precios del SCOP entre el 13 de marzo de 2020 y el 24 de septiembre de 2026 para cinco "
     "combustibles (diésel B5 S-50, gasolina regular, gasolina premium, GLP envasado de 10 kg y GLP a granel), hemos "
     "construido un método que ordena diariamente los establecimientos de venta al público según la probabilidad de que su "
     "precio cambie en los cinco días siguientes, como apoyo a la planificación de actividades de supervisión. El método "
     "no determina incumplimientos ni reemplaza los criterios técnicos de la entidad.", alinear=J)
p(c, "En el tratamiento de la información, el código de registro de cada establecimiento se reemplazó por un seudónimo; "
     "el RUC, la razón social y la dirección no se utilizan en el análisis; no se identifica a personas naturales; y los "
     "resultados se reportan de forma agregada.", alinear=J)
p(c, "Si bien la información es de acceso público, por buenas prácticas de investigación y respeto a la institución que la "
     "administra, solicitamos a usted lo siguiente:", alinear=J)
lista(c, [
    "**Conformidad para el uso de la información con fines exclusivamente académicos** (tesis de pregrado y un artículo "
    "científico), y que se nos indique, de ser el caso: a) si la conformidad comprende toda la información o solo "
    "determinados campos, productos o periodos; b) si los resultados pueden publicarse y con qué nivel de detalle "
    "(nacional, regional o por establecimiento); c) la forma en que debe citarse a Osinergmin como fuente; d) si en el "
    "prototipo de demostración pueden mostrarse las razones sociales de personas jurídicas o debe usarse únicamente el "
    "seudónimo; y e) cualquier otra condición que considere pertinente.",
    "**La designación de uno o más especialistas como expertos validadores**, para participar en una validación por juicio "
    "de expertos (método Delphi) de dos rondas, con un cuestionario breve de aproximadamente 30 minutos cada una. Sus "
    "respuestas se analizarán de forma agregada y anónima.",
    "**Una reunión de aproximadamente 45 minutos**, virtual o presencial, para presentar el método y su prototipo de "
    "consulta, y recoger sus observaciones.",
], numerada=True)
p(c, "Por nuestra parte, nos comprometemos a utilizar la información sin fines comerciales, a no redistribuir las bases de "
     "datos, a respetar las condiciones que se nos indiquen, a precisar en las publicaciones que las opiniones y resultados "
     "son de exclusiva responsabilidad de los autores y no representan la posición de Osinergmin, y a entregar una copia del "
     "documento final.", alinear=J, antes=4)
p(c, "Adjuntamos un resumen ejecutivo del proyecto y [la carta de presentación de la universidad o del asesor]. Para "
     "cualquier coordinación, puede comunicarse con nosotros al correo [correo] o al teléfono [teléfono].", alinear=J)
p(c, "Agradeciendo de antemano la atención a la presente, quedamos de usted.", alinear=J, despues=6)
p(c, "Atentamente,", despues=40)

firmas = c.add_table(rows=2, cols=3)
nombres = ["David Edison Guerra Ordoñez", "Richard [ ] Yanqui Inocente", "Kevin Javier Tello Rivas"]
for j, nmb in enumerate(nombres):
    for i, txt in enumerate(["______________________", f"{nmb}\nDNI N.° [ ]"]):
        cel = firmas.cell(i, j); cel.text = ""
        par = cel.paragraphs[0]; par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run(txt); r.font.size = Pt(9.5)
p(c, "", despues=30)
p(c, "______________________", alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=0)
p(c, "V.° B.° [Nombres y apellidos del asesor]", alinear=WD_ALIGN_PARAGRAPH.CENTER, tam=9.5, despues=0)
p(c, "Asesor de tesis – Universidad Peruana Unión", alinear=WD_ALIGN_PARAGRAPH.CENTER, tam=9.5, despues=14)
p(c, "Adjuntos: (1) Resumen ejecutivo del proyecto (1 página). (2) [Carta de presentación de la universidad o del asesor].",
  tam=9.5, cursiva=True)
c.save(os.path.join(AQUI, "carta_solicitud_osinergmin.docx"))

# ================================================================== GUÍA
g = base(2.2)
par = p(g, "Guía para presentar el proyecto a Osinergmin", negrita=True, tam=17, despues=2)
for r in par.runs:
    r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
p(g, "Uso interno del equipo de tesis · octubre de 2026", tam=9.5, cursiva=True, despues=10)

titulo(g, "1. El mensaje en una frase")
p(g, "«Con los precios que Osinergmin ya recibe y publica, proponemos un **método para ordenar cada día a qué "
     "establecimientos conviene mirar primero**, porque su precio tiene más probabilidad de cambiar en los próximos 5 días. "
     "Queremos saber si les sería útil y qué le falta.»", alinear=J)
p(g, "Se presenta como un **método de priorización a validar con ellos**, no como la solución a un problema grave que "
     "nosotros descubrimos. Así el problema no depende de afirmar cosas sobre Osinergmin que no podemos probar.", alinear=J)

titulo(g, "2. Qué podemos afirmar y qué no")
p(g, "Revisamos los documentos que tenemos de Osinergmin. Esto es lo que dicen y lo que no:", alinear=J)
tabla(g, [
    ["Afirmación", "Respaldo", "¿Se puede decir?"],
    ["Los establecimientos deben reportar sus precios al SCOP y Osinergmin los publica.",
     "Informe GSE/DSR-1875-2026, que cita el procedimiento PRICE (RCD N.° 256-2021-OS/CD, art. 17).", "**Sí.** Es la base más sólida."],
    ["Osinergmin estudió el fenómeno «cohetes y plumas» (los precios suben rápido y bajan lento).",
     "Reporte de Análisis Económico Sectorial de Hidrocarburos, marzo de 2016; Documentos de Trabajo 39 (2017) y 49 (2020) de la GPAE.",
     "**Sí, como contexto.** Aclarar que nuestro modelo predice si el precio cambia, no mide esa asimetría."],
    ["Osinergmin pide un análisis más riguroso.",
     "El DT 39 concluye que hay evidencia preliminar de asimetría en el GLP de Lima y Lambayeque y que es necesario continuar la investigación.",
     "**Solo así:** «sus propios estudios recomiendan seguir investigando la dinámica de precios»."],
    ["Osinergmin quiere priorizar establecimientos para supervisar.",
     "No lo encontramos en esos documentos.", "**No.** Mejor preguntárselo en la reunión."],
    ["El modelo detecta infractores o precios mal reportados.",
     "No: el modelo predice cambios de precio, no incumplimientos.", "**No.**"],
    ["Que un cambio de precio sea un momento de riesgo para el consumidor.",
     "Es una hipótesis de la tesis; no se comprobó con datos de supervisión.", "**Como hipótesis**, y pedir su opinión."],
], [5.2, 7.3, 4.1])

titulo(g, "3. Cómo plantear el problema sin exagerar")
p(g, "El problema se sostiene con los datos, sin necesidad de atribuirle nada a Osinergmin:", alinear=J)
lista(g, [
    "Hay **24 771 establecimientos** que reportan precios de 5 combustibles.",
    "El precio **no cambia en el 95,35 % de los días**: los cambios son pocos y difíciles de anticipar a simple vista.",
    "Ninguna institución puede mirar todo cada día, así que **alguien tiene que decidir por dónde empezar**.",
    "Proponemos una forma ordenada, medible y transparente de hacerlo, que explica cada caso con razones (SHAP).",
])

titulo(g, "4. Cifras que se pueden mostrar (prueba de enero a setiembre de 2026)")
tabla(g, [
    ["Qué se mide", "Modelo", "Regla simple", "Cómo decirlo"],
    ["Orden de toda la lista (AUC diario)", "0,746", "0,725 y 0,704", "Ordena mejor en los 9 meses."],
    ["Aciertos en los 10 primeros de cada región", "66 %", "64 %", "Ligera ventaja en la escala regional."],
    ["Aciertos en los 100 primeros del país", "78 %", "79 %", "**Empate.** Decirlo así; fue por el choque de precios de marzo."],
    ["Establecimientos nunca vistos (100 primeros)", "74 %", "—", "Funciona con grifos que no usó para aprender."],
    ["Elegir al azar", "≈ 33 %", "", "Referencia para ver cuánto aporta ordenar."],
], [5.6, 1.8, 2.4, 6.8])

titulo(g, "5. Estructura de la reunión (45 minutos)")
tabla(g, [
    ["Minutos", "Qué hacer"],
    ["0–5", "Presentación del equipo y agradecimiento por la información (citar el Informe GSE/DSR-1875-2026)."],
    ["5–15", "Problema y método: las 4 ideas de la sección 3 y un diagrama. Sin fórmulas."],
    ["15–30", "Demostración del sistema: lista del día por región, razones de un caso, comprobación contra el SCOP. Usar el guion 07 del kit Delphi."],
    ["30–40", "**Escuchar**: hacer las preguntas de la sección 6."],
    ["40–45", "Próximos pasos: conformidad de uso, especialistas validadores y fecha de la ronda 1 del Delphi."],
], [2.0, 14.6])

titulo(g, "6. Preguntas para hacerles")
lista(g, [
    "¿Cómo deciden hoy qué establecimientos supervisar primero en temas de precios?",
    "¿Qué información usan para esa decisión?",
    "¿Les serviría una lista diaria por región? ¿De qué largo?",
    "¿Qué le faltaría al método para que lo consideren útil?",
    "¿Existe información de resultados de supervisión que podría usarse, de forma anónima, para comprobar la hipótesis de la tesis?",
])

titulo(g, "7. Preguntas difíciles y respuestas honestas")
tabla(g, [
    ["Si preguntan…", "Responder…"],
    ["¿Por qué un cambio de precio indica riesgo?",
     "Es la hipótesis de trabajo: al cambiar el precio puede haber diferencia entre lo registrado y lo cobrado. No la comprobamos con datos de supervisión porque no son públicos; por eso pedimos su opinión."],
    ["En los 100 primeros, el modelo empata con una regla simple.",
     "Es cierto. Gana en el orden de toda la lista, en las listas regionales y funciona con establecimientos nuevos; además explica cada caso."],
    ["¿Esto reemplaza al supervisor?", "No. Solo ordena; la decisión siempre es del especialista."],
    ["¿Qué pasa con los datos de las empresas?", "Se usan con seudónimo, sin RUC ni dirección, y los resultados se publican de forma agregada."],
], [5.2, 11.4])

titulo(g, "8. Qué llevar")
lista(g, [
    "Carta de solicitud firmada y resumen ejecutivo de 1 página.",
    "Sistema funcionando en una laptop (iniciar_sistema.bat) y, como respaldo, el dashboard estático.",
    "Diagrama metodológico y ficha del kit Delphi (02).",
    "Copia del Informe GSE/DSR-1875-2026.",
])
g.save(os.path.join(AQUI, "guia_reunion_osinergmin.docx"))
print("ok")
